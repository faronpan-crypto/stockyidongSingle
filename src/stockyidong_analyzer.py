#!/usr/bin/env python3
"""
stockyidong_analyzer.py
读取每日导出的 stockyidong 文件，按游资共振逻辑排序，生成 Word 分析报告。

用法：
  python3 stockyidong_analyzer.py                    # 自动读取最新文件
  python3 stockyidong_analyzer.py --input /path/to/stockyidong_20260506_102050.txt
"""

import glob
import os
import re
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

# ========== 配置 ==========
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_INPUT_DIR = SCRIPT_DIR
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "report")

# 龙头权重说明
WEIGHT_LABELS = {
    "龙头股": "A（最强主线龙头）",
    "15Min": "B（短线热点）",
    "同花顺": "C（市场资金共振）",
}

# ========== 解析 ==========
def parse_stock_entry(text):
    """从 '圣阳股份(002580)' 提取 (name, code)"""
    m = re.match(r"^(.+?)\((\d{6})\)$", text.strip())
    if m:
        return m.group(1).strip(), m.group(2)
    return None, None

def parse_section(block):
    """解析一个标签页块，返回 [(name, code), ...]"""
    stocks = []
    # 找 "【龙头股】" 之后的内容
    m = re.search(r"【.+?】\s*(.+?)(?=\n【|\n+$)", block, re.DOTALL)
    if m:
        raw = m.group(1).strip()
        for seg in raw.split("、"):
            seg = seg.strip()
            if not seg or seg in ("（无）", "(无)"):
                continue
            name, code = parse_stock_entry(seg)
            if name and code:
                stocks.append({"name": name, "code": code})
    return stocks

def load_latest_file(indir):
    """找最新的 stockyidong_*.txt"""
    pattern = os.path.join(indir, "stockyidong_*.txt")
    files = glob.glob(pattern)
    if not files:
        raise FileNotFoundError(f"未找到文件: {pattern}")
    latest = max(files, key=os.path.getmtime)
    print(f"[INFO] 读取文件: {latest}")
    with open(latest, "r", encoding="utf-8") as f:
        return latest, f.read()

def parse_content(content):
    """按标签页解析文件内容"""
    sections = {}
    for label in ["龙头股", "15Min", "同花顺"]:
        # 切出每个标签页的块
        pattern = rf"【{label}】.*?(?=【|\Z)"
        m = re.search(pattern, content, re.DOTALL)
        if m:
            sections[label] = parse_section(m.group(0))
        else:
            sections[label] = []
    return sections

# ========== 排序逻辑 ==========
def build_score_table(sections):
    """计算每只股票的共振分并排序"""
    code_info = {}  # code -> {name, tabs: []}
    tab_order = ["龙头股", "15Min", "同花顺"]

    for tab, stocks in sections.items():
        for s in stocks:
            code = s["code"]
            if code not in code_info:
                code_info[code] = {"name": s["name"], "tabs": [], "tab_order": []}
            if tab not in code_info[code]["tabs"]:
                code_info[code]["tabs"].append(tab)
                code_info[code]["tab_order"].append(tab)

    scored = []
    for code, info in code_info.items():
        n = len(info["tabs"])
        # 加权：龙头=3分，15Min=2分，同花顺=1分
        score = 0
        for t in info["tabs"]:
            if t == "龙头股":
                score += 3
            elif t == "15Min":
                score += 2
            elif t == "同花顺":
                score += 1
        scored.append({
            "code": code,
            "name": info["name"],
            "tabs": info["tabs"],
            "tab_count": n,
            "score": score,
        })

    # 按共振分降序，同分按出现标签页数降序，再按龙头→15Min→同花顺顺序
    def sort_key(s):
        tab_idx = { "龙头股": 0, "15Min": 1, "同花顺": 2 }
        first_tab = s["tabs"][0] if s["tabs"] else "同花顺"
        return (-s["score"], -s["tab_count"], tab_idx.get(first_tab, 2))
    scored.sort(key=sort_key)
    return scored

# ========== Word 生成 ==========
def make_word(scored, sections, now_str, out_path):
    doc = Document()

    # 页面边距
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin  = Cm(2.5)
        section.right_margin = Cm(2.5)

    # 标题
    title = doc.add_heading("斯东克 · 每日股票共振分析", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.runs[0]
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    doc.add_paragraph(now_str).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 总览
    total = len(scored)
    doc.add_paragraph(
        f"📊 总覆盖：{total} 只 | 龙头股 {len(sections['龙头股'])} 只 | "
        f"15Min {len(sections['15Min'])} 只 | 同花顺 {len(sections['同花顺'])} 只"
    )

    # ===== 共振排行 =====
    h = doc.add_heading("一、共振排行（按跨标签共振强度排序）", level=2)
    h.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    doc.add_paragraph(
        "说明：每只股票按「龙头股(×3) + 15Min(×2) + 同花顺(×1)」加权计分，"
        "分数越高说明在多个维度同时获得游资/量化资金认可。"
    )

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "排名"
    hdr[1].text = "股票名称"
    hdr[2].text = "代码"
    hdr[3].text = "共振标签"
    for cell in hdr:
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(10)

    for rank, s in enumerate(scored[:30], 1):
        row = table.add_row().cells
        row[0].text = str(rank)
        row[1].text = s["name"]
        row[2].text = s["code"]
        row[3].text = " / ".join(s["tabs"])

        # 前3名高亮
        if rank <= 3:
            for cell in row:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
                        run.bold = True

    # ===== 各标签页详细列表 =====
    doc.add_paragraph()
    h2 = doc.add_heading("二、各标签页明细", level=2)
    h2.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    for tab in ["龙头股", "15Min", "同花顺"]:
        stocks = sections[tab]
        doc.add_heading(f"【{tab}】{WEIGHT_LABELS[tab]}", level=3)
        names = "、".join(f"{s['name']}({s['code']})" for s in stocks)
        doc.add_paragraph(names if names else "（无）")

    doc.save(out_path)
    return out_path

# ========== 主流程 ==========
def main():
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", "-i", default=None)
    parser.add_argument("--out", "-o", default=None)
    args = parser.parse_args()

    # 输入
    if args.input:
        with open(args.input, "r", encoding="utf-8") as f:
            content = f.read()
        input_path = args.input
    else:
        input_path, content = load_latest_file(DEFAULT_INPUT_DIR)

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    sections = parse_content(content)
    scored = build_score_table(sections)

    # 输出
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    if args.out:
        out_path = args.out
    else:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_path = os.path.join(OUTPUT_DIR, f"stockyidong_report_{ts}.docx")

    make_word(scored, sections, now_str, out_path)
    print(f"[OK] Word 报告已生成: {out_path}")
    return out_path, scored[:5]  # 返回路径 + Top5

if __name__ == "__main__":
    out_path, top5 = main()
    print("\n🔥 Top5 共振股票:")
    for i, s in enumerate(top5, 1):
        print(f"  {i}. {s['name']}({s['code']}) — {s['tabs']} ({s['score']}分)")
